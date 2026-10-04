#!/usr/bin/env python3
"""
Professional Minimalist Executive Slide Deck (16:9 Widescreen)
Style: Modern Minimalist Light (McKinsey / Stripe / Tech Analytics Style)
"""

import os
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from report_metrics import load_metrics, money, money_m, pct

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PRES_DIR = os.path.join(BASE_DIR, "presentations")
VIZ_DIR = os.path.join(BASE_DIR, "visualizations")
os.makedirs(PRES_DIR, exist_ok=True)

# -----------------------------------------------------------------------------
# MINIMALIST LIGHT PALETTE
# -----------------------------------------------------------------------------
C_CANVAS = RGBColor(248, 250, 252)     # Slate 50 (#F8FAFC)
C_CARD = RGBColor(255, 255, 255)       # Pure White (#FFFFFF)
C_CARD_MUTED = RGBColor(241, 245, 249) # Slate 100 (#F1F5F9)
C_BORDER = RGBColor(226, 232, 240)     # Slate 200 (#E2E8F0)
C_BORDER_DARK = RGBColor(203, 213, 225)# Slate 300 (#CBD5E1)

C_PRIMARY = RGBColor(15, 23, 42)       # Slate 900 / Charcoal (#0F172A)
C_BODY = RGBColor(51, 65, 85)          # Slate 700 (#334155)
C_MUTED = RGBColor(100, 116, 139)      # Slate 500 (#64748B)

C_BLUE = RGBColor(37, 99, 235)         # Royal Blue 600 (#2563EB)
C_BLUE_BG = RGBColor(239, 246, 255)    # Blue 50 (#EFF6FF)
C_AMBER = RGBColor(217, 119, 6)        # Amber 600 (#D97706)
C_EMERALD = RGBColor(16, 185, 129)     # Emerald 500 (#10B981)
C_ROSE = RGBColor(225, 29, 72)         # Rose 600 (#E11D48)

def set_shape_style(shape, fill_color=C_CARD, line_color=C_BORDER, line_width=1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()

def add_header(slide, title_text, category_text="EXECUTIVE DATA ANALYTICS CASE STUDY"):
    # Category Pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(3.6), Inches(0.28))
    set_shape_style(pill, fill_color=C_BLUE_BG, line_color=RGBColor(191, 219, 254), line_width=0.8)
    ptf = pill.text_frame
    ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = 0
    p = ptf.paragraphs[0]
    p.text = category_text.upper()
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    # Title Text
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.55))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p2 = tf.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = C_PRIMARY

def generate_slides():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    m = load_metrics()
    seg = m["seg"]
    champ, loyal, potential, risk = (seg.loc[k] for k in ("Champions", "Loyal Customers", "Potential Loyalists", "At Risk"))
    cm1 = m["cohort_m1"]
    nov10, dec10 = cm1[pd.Timestamp("2010-11-01")], cm1[pd.Timestamp("2010-12-01")]

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        set_shape_style(bg, fill_color=C_CANVAS, line_color=None)
        return bg

    # =========================================================================
    # SLIDE 1: COVER / HERO (Minimalist, Crisp, High-Impact)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    hero_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    set_shape_style(hero_card, fill_color=C_CARD, line_color=C_BORDER, line_width=1)

    tb1 = s1.shapes.add_textbox(Inches(1.3), Inches(1.3), Inches(10.7), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0

    p_pill = tf1.paragraphs[0]
    p_pill.text = "PORTFOLIO CASE STUDY • RETAIL & E-COMMERCE ANALYTICS"
    p_pill.font.size = Pt(10)
    p_pill.font.bold = True
    p_pill.font.color.rgb = C_BLUE
    p_pill.space_after = Pt(12)

    p_h1 = tf1.add_paragraph()
    p_h1.text = "Customer Segmentation & Cohort Retention Optimization"
    p_h1.font.size = Pt(30)
    p_h1.font.bold = True
    p_h1.font.color.rgb = C_PRIMARY
    p_h1.space_after = Pt(10)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Tackling a 25% Surge in Customer Acquisition Costs (CAC) via Behavioral RFM Modeling, Month-1 Churn Diagnostics, and Strategic Lifecycle Marketing."
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = C_MUTED
    p_sub.space_after = Pt(20)

    p_meta = tf1.add_paragraph()
    p_meta.text = "Author: Novaldi Ramadhan Waluyo (Data Analyst)   |   Stakeholders: CMO, CRM Lead, Product Growth Manager"
    p_meta.font.size = Pt(10)
    p_meta.font.bold = True
    p_meta.font.color.rgb = C_BODY

    # 4 Bottom Stat Chips
    stat_chips = [
        ("TOTAL REVENUE", money_m(m["revenue"]), C_BLUE),
        ("CLEANED TRANSACTIONS", f"{m['clean_rows'] / 1e3:.1f}K", C_EMERALD),
        ("UNIQUE CUSTOMERS", f"{m['customers']:,}", C_PRIMARY),
        ("NOT BACK IN MONTH 1", pct(m["m1_churn"]), C_ROSE)
    ]
    chip_w = Inches(2.55)
    chip_h = Inches(1.15)
    chip_y = Inches(5.05)
    for i, (lbl, val, col) in enumerate(stat_chips):
        cx = Inches(1.3 + i * 2.72)
        chip = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, chip_y, chip_w, chip_h)
        set_shape_style(chip, fill_color=C_CARD_MUTED, line_color=C_BORDER, line_width=1)
        ctf = chip.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ctf.margin_left = ctf.margin_right = ctf.margin_top = ctf.margin_bottom = 0
        
        cp1 = ctf.paragraphs[0]
        cp1.text = lbl
        cp1.font.size = Pt(8.5)
        cp1.font.bold = True
        cp1.font.color.rgb = C_MUTED
        cp1.alignment = PP_ALIGN.CENTER
        
        cp2 = ctf.add_paragraph()
        cp2.text = val
        cp2.font.size = Pt(18)
        cp2.font.bold = True
        cp2.font.color.rgb = col
        cp2.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 2: SCORECARD & STRATEGIC PROBLEM STATEMENT
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Executive Scorecard & Strategic Problem Statement")

    # 4 Top KPI Cards
    top_metrics = [
        ("TOTAL REVENUE", money_m(m["revenue"]), f"{m['orders']:,} orders, Dec 2009 - Dec 2011", C_BLUE),
        ("REGISTERED CUSTOMERS", f"{m['customers']:,}", f"{m['countries']} countries", C_PRIMARY),
        ("MONTH-1 RETENTION", pct(m["m1"]), "Steepest drop-off in the lifecycle", C_ROSE),
        ("CHAMPIONS + LOYAL REVENUE", pct(m["core_pct_revenue"]), f"From {pct(m['core_pct_customers'])} of customers", C_EMERALD)
    ]
    kw = Inches(2.78)
    kh = Inches(1.2)
    ky = Inches(1.35)
    for i, (title, val, sub, col) in enumerate(top_metrics):
        kx = Inches(0.8 + i * 2.98)
        kcard = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, kx, ky, kw, kh)
        set_shape_style(kcard, fill_color=C_CARD, line_color=C_BORDER, line_width=1)
        ktf = kcard.text_frame
        ktf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ktf.margin_left = ktf.margin_right = Inches(0.18)
        ktf.margin_top = ktf.margin_bottom = 0
        
        p1 = ktf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(8)
        p1.font.bold = True
        p1.font.color.rgb = C_MUTED
        
        p2 = ktf.add_paragraph()
        p2.text = val
        p2.font.size = Pt(19)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.space_before = Pt(2)
        
        p3 = ktf.add_paragraph()
        p3.text = sub
        p3.font.size = Pt(7.5)
        p3.font.color.rgb = C_MUTED
        p3.space_before = Pt(2)

    # 2 Comparison Panels Below
    panel_y = Inches(2.75)
    panel_h = Inches(4.25)
    panel_w = Inches(5.72)

    # Left Panel: Problem Diagnosis
    lp = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), panel_y, panel_w, panel_h)
    set_shape_style(lp, fill_color=C_CARD, line_color=C_BORDER, line_width=1)
    lptf = lp.text_frame
    lptf.word_wrap = True
    lptf.margin_left = lptf.margin_right = Inches(0.25)
    lptf.margin_top = Inches(0.2)
    lptf.margin_bottom = Inches(0.15)

    p = lptf.paragraphs[0]
    p.text = "BUSINESS PROBLEM & VULNERABILITY DIAGNOSIS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_ROSE
    p.space_after = Pt(8)

    diag_points = [
        ("Surging Acquisition Costs (+25% CAC):", "Rising ad channel costs mean acquiring one-off buyers severely damages long-term profitability."),
        ("Inefficient Blanket Discounting:", "Mass promo discounts eroded profit margins without driving repeat loyalty or sustainable habit formation."),
        (f"The Month-1 Cliff ({pct(m['m1_churn'])} not back):", "Roughly three in four new customers do not buy again in the month after their first order."),
        ("Silent VIP Value Decay:", f"{int(risk.customers):,} once-valuable customers ({money_m(risk.revenue)} historical revenue) have drifted into At Risk.")
    ]
    for dt, dd in diag_points:
        pt = lptf.add_paragraph()
        pt.text = f"• {dt} "
        pt.font.bold = True
        pt.font.size = Pt(9.5)
        pt.font.color.rgb = C_PRIMARY
        pt.space_before = Pt(6)
        
        pdesc = lptf.add_paragraph()
        pdesc.text = f"  {dd}"
        pdesc.font.size = Pt(8.8)
        pdesc.font.color.rgb = C_BODY
        pdesc.space_before = Pt(1)

    # Right Panel: Strategic Objectives
    rp = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.81), panel_y, panel_w, panel_h)
    set_shape_style(rp, fill_color=C_CARD, line_color=C_BORDER, line_width=1)
    rptf = rp.text_frame
    rptf.word_wrap = True
    rptf.margin_left = rptf.margin_right = Inches(0.25)
    rptf.margin_top = Inches(0.2)
    rptf.margin_bottom = Inches(0.15)

    p = rptf.paragraphs[0]
    p.text = "STRATEGIC OBJECTIVES & REMEDIATION GOALS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_after = Pt(8)

    obj_points = [
        ("Reallocate 20% Budget to Retention:", "Shift paid media budget from low-intent cold ads to high-ROI automated post-purchase flows."),
        ("Eliminate Blanket Promo Discounts:", "Preserve margin by replacing general discounts with VIP perks for Champions and win-back offers for At-Risk."),
        ("Bridge the Month-1 Gap:", "Send new buyers personalised category recommendations within 14 days of their first order."),
        ("Track Retention in Power BI:", "Give growth teams a dashboard for month-1 retention, cohort curves and segment shifts.")
    ]
    for ot, od in obj_points:
        pt = rptf.add_paragraph()
        pt.text = f"• {ot} "
        pt.font.bold = True
        pt.font.size = Pt(9.5)
        pt.font.color.rgb = C_PRIMARY
        pt.space_before = Pt(6)
        
        pdesc = rptf.add_paragraph()
        pdesc.text = f"  {od}"
        pdesc.font.size = Pt(8.8)
        pdesc.font.color.rgb = C_BODY
        pdesc.space_before = Pt(1)

    # =========================================================================
    # SLIDE 3: COHORT RETENTION ANALYSIS (Snug Fit Chart + 3 Side Cards)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Cohort Retention Analysis: Mapping the Month-1 Churn Cliff")

    # Left Card Container for Chart (Exact Dimensions)
    chart_box_w = Inches(7.5)
    chart_box_h = Inches(5.6)
    c_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), chart_box_w, chart_box_h)
    set_shape_style(c_box, fill_color=C_CARD, line_color=C_BORDER, line_width=1)

    # Embed Chart perfectly centered inside box
    img_cohort = os.path.join(VIZ_DIR, "01_cohort_retention_heatmap.png")
    if os.path.exists(img_cohort):
        # 7.2" width fits comfortably in 7.5" card with 0.15" margins
        s3.shapes.add_picture(img_cohort, Inches(0.95), Inches(1.5), width=Inches(7.2))

    # Right Column: 3 Snug Structured Insight Cards
    side_w = Inches(4.03)
    side_h = Inches(1.76)
    side_x = Inches(8.5)

    c_cards = [
        (f"THE MONTH-1 CLIFF ({pct(m['m1'])} RETAINED)",
         f"Only {pct(m['m1'])} of new customers buy again in the month after their first order (customer-weighted across all cohorts).",
         C_ROSE),
        ("FLAT AFTER THE CLIFF",
         f"Retention then holds between {pct(m['retention_min_m1_m12'])} and {pct(m['retention_max_m1_m12'])} from month 1 to month 12, so the second purchase is the main lever.",
         C_BLUE),
        ("HOLIDAY COHORTS RETAIN WORST",
         f"Customers first acquired in Nov 2010 ({pct(nov10)}) and Dec 2010 ({pct(dec10)}) came back far less in month 1. Follow up within 30 days.",
         C_AMBER)
    ]
    for i, (title, desc, col) in enumerate(c_cards):
        cy = Inches(1.35 + i * 1.92)
        sc = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, side_x, cy, side_w, side_h)
        set_shape_style(sc, fill_color=C_CARD, line_color=C_BORDER, line_width=1)
        sctf = sc.text_frame
        sctf.word_wrap = True
        sctf.margin_left = sctf.margin_right = Inches(0.18)
        sctf.margin_top = Inches(0.15)
        sctf.margin_bottom = Inches(0.1)

        p1 = sctf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_after = Pt(4)

        p2 = sctf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(8.6)
        p2.font.color.rgb = C_BODY
        p2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 4: RFM BEHAVIORAL SEGMENTATION (Chart + 3 Segment Cards)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "RFM Behavioral Segmentation: Pareto Concentration & Tiers")

    rfm_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), chart_box_w, chart_box_h)
    set_shape_style(rfm_box, fill_color=C_CARD, line_color=C_BORDER, line_width=1)

    img_rfm = os.path.join(VIZ_DIR, "02_rfm_revenue_vs_customers.png")
    if os.path.exists(img_rfm):
        s4.shapes.add_picture(img_rfm, Inches(0.95), Inches(1.5), width=Inches(7.2))

    rfm_cards = [
        (f"CHAMPIONS ({pct(champ.pct_revenue)} REVENUE | {pct(champ.pct_customers)} CUSTOMERS)",
         f"{int(champ.customers):,} customers generate {money_m(champ.revenue)} with {champ.orders:.1f} orders on average ({money(champ.spend)} spend). Action: VIP perks and early access, no discounts.",
         C_EMERALD),
        (f"LOYAL & POTENTIAL ({pct(loyal.pct_revenue + potential.pct_revenue)} REVENUE)",
         f"{int(loyal.customers + potential.customers):,} customers contributing {money_m(loyal.revenue + potential.revenue)}. Action: loyalty milestones and cross-category recommendations.",
         C_BLUE),
        (f"AT RISK ({pct(risk.pct_revenue)} REVENUE)",
         f"{int(risk.customers):,} customers averaging {money(risk.spend)} spend, last order ~{risk.recency:.0f} days ago ({money_m(risk.revenue)} historical revenue). Action: targeted win-back offers.",
         C_AMBER)
    ]
    for i, (title, desc, col) in enumerate(rfm_cards):
        cy = Inches(1.35 + i * 1.92)
        sc = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, side_x, cy, side_w, side_h)
        set_shape_style(sc, fill_color=C_CARD, line_color=C_BORDER, line_width=1)
        sctf = sc.text_frame
        sctf.word_wrap = True
        sctf.margin_left = sctf.margin_right = Inches(0.18)
        sctf.margin_top = Inches(0.15)
        sctf.margin_bottom = Inches(0.1)

        p1 = sctf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_after = Pt(4)

        p2 = sctf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(8.6)
        p2.font.color.rgb = C_BODY
        p2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 5: REVENUE & ORDER TRAJECTORY (Chart + 3 Temporal Cards)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Monthly Revenue & Order Volume Dynamics (2009 - 2011)")

    gro_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), chart_box_w, chart_box_h)
    set_shape_style(gro_box, fill_color=C_CARD, line_color=C_BORDER, line_width=1)

    img_growth = os.path.join(VIZ_DIR, "03_monthly_revenue_growth.png")
    if os.path.exists(img_growth):
        s5.shapes.add_picture(img_growth, Inches(0.95), Inches(1.5), width=Inches(7.2))

    gro_cards = [
        ("Q4 PEAK (~£1.2M / MONTH)",
         "Revenue climbs from September and peaks in November in both years, roughly twice the Q1 monthly level.",
         C_BLUE),
        ("POST-HOLIDAY RE-ENGAGEMENT",
         "Holiday first-time buyers retain worst, so a January re-engagement flow keeps them from becoming Hibernating.",
         C_AMBER),
        ("DECEMBER 2011 IS PARTIAL",
         f"The data ends on 9 Dec 2011, so the final drop is a cut-off, not a decline. Average order value over the period: {money(m['aov'])}.",
         C_EMERALD)
    ]
    for i, (title, desc, col) in enumerate(gro_cards):
        cy = Inches(1.35 + i * 1.92)
        sc = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, side_x, cy, side_w, side_h)
        set_shape_style(sc, fill_color=C_CARD, line_color=C_BORDER, line_width=1)
        sctf = sc.text_frame
        sctf.word_wrap = True
        sctf.margin_left = sctf.margin_right = Inches(0.18)
        sctf.margin_top = Inches(0.15)
        sctf.margin_bottom = Inches(0.1)

        p1 = sctf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_after = Pt(4)

        p2 = sctf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(8.6)
        p2.font.color.rgb = C_BODY
        p2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 6: 3-PILLAR STAKEHOLDER STRATEGIC PLAYBOOK
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "3-Pillar Stakeholder Strategic Playbook")

    pillars = [
        ("HEAD OF MARKETING / CMO", C_BLUE, [
            ("Reallocate 20% Ad Budget", "Shift media spend from low-intent cold search/social to automated lifecycle retention flows."),
            ("Eliminate Public Discounting", "Stop blanket promo codes; protect product prestige and gross profit margins."),
            ("Tiered Promo Strategy", "Provide exclusive VIP previews to Champions; reserve financial incentives strictly for At-Risk win-back."),
            ("Measure with Holdout Tests", "Run each change against a control group to measure its real effect on repeat revenue and CAC.")
        ]),
        ("CRM & RETENTION LEAD", C_AMBER, [
            ("14-Day Post-Purchase Trigger", "Deploy automated onboarding emails offering cross-category complements based on 1st order SKU."),
            ("60-Day Inactivity Warning", "Trigger automated high-urgency win-back sequence before customer recency exceeds 180 days."),
            ("Multi-Channel Orchestration", "Deliver high-touch SMS/WhatsApp for urgent At-Risk accounts and weekly curated digests for Loyalists."),
            ("Dynamic RFM Scoring", "Recalculate customer segments weekly to detect segment transitions in real time.")
        ]),
        ("PRODUCT & GROWTH MANAGER", C_EMERALD, [
            ("Tiered Loyalty Program", "Implement Bronze, Silver, Gold, Platinum tiers with tangible milestone rewards on 3rd & 5th orders."),
            ("1-Click Frictionless Reorder", "Add instant reorder buttons on order history pages for high-frequency consumable items."),
            ("Power BI Retention Dashboard", "Monitor month-1 retention, cohort curves and segment shifts in the three-page Power BI report."),
            ("Cross-Sell Recommendation Engine", "Deploy intelligent recommendation widgets during checkout based on product affinity rules.")
        ])
    ]
    col_w = Inches(3.78)
    col_h = Inches(5.6)
    for i, (col_title, col_accent, items) in enumerate(pillars):
        cx = Inches(0.8 + i * 3.97)
        ccard = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.35), col_w, col_h)
        set_shape_style(ccard, fill_color=C_CARD, line_color=C_BORDER, line_width=1)
        cctf = ccard.text_frame
        cctf.word_wrap = True
        cctf.margin_left = cctf.margin_right = Inches(0.2)
        cctf.margin_top = Inches(0.2)
        cctf.margin_bottom = Inches(0.15)

        p = cctf.paragraphs[0]
        p.text = col_title
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = col_accent
        p.space_after = Pt(8)

        for item_title, item_desc in items:
            ip = cctf.add_paragraph()
            ip.text = f"• {item_title}"
            ip.font.size = Pt(9.2)
            ip.font.bold = True
            ip.font.color.rgb = C_PRIMARY
            ip.space_before = Pt(6)

            idp = cctf.add_paragraph()
            idp.text = f"  {item_desc}"
            idp.font.size = Pt(8.4)
            idp.font.color.rgb = C_BODY
            idp.space_before = Pt(1)

    # =========================================================================
    # SLIDE 7: TECHNICAL DATA ARCHITECTURE & AUTOMATION
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Technical Architecture & Automated Data Infrastructure")

    tech_steps = [
        ("1. DATA EXTRACTION & CLEANING", C_BLUE, [
            f"Processed {m['raw_rows'] / 1e6:.2f}M raw Excel lines across 2 yearly sheets.",
            "Removed cancellations ('C'), returns, zero prices and guest checkouts.",
            f"Kept {m['clean_rows']:,} line items, {money_m(m['revenue'])} revenue."
        ]),
        ("2. SQL & ANALYTICAL MODELING", C_PRIMARY, [
            "Built SQLite schema (`online_retail_analytics.db`).",
            "Engineered RFM scores (`NTILE(5)`) and Cohort Index offsets.",
            "Created indexed views: `v_rfm_summary`, `v_monthly_sales_trend`."
        ]),
        ("3. POWER BI & TABLEAU", C_EMERALD, [
            "Power BI star schema, DAX measures and theme in `/powerbi/`.",
            "Three-page dashboard: Overview, Cohort Retention, RFM Segments.",
            "Tableau data source (`.tds`) for a live SQLite connection via ODBC."
        ]),
        ("4. REPRODUCIBLE ARTIFACTS", C_AMBER, [
            "Production SQL scripts in `/sql/` directory.",
            "All figures computed in `src/report_metrics.py`.",
            "High-resolution 300 DPI visualizations & detailed PDF briefing."
        ])
    ]
    tcard_w = Inches(2.78)
    tcard_h = Inches(5.6)
    for i, (st_title, st_col, st_points) in enumerate(tech_steps):
        tx = Inches(0.8 + i * 2.98)
        tcard = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, Inches(1.35), tcard_w, tcard_h)
        set_shape_style(tcard, fill_color=C_CARD, line_color=C_BORDER, line_width=1)
        tctf = tcard.text_frame
        tctf.word_wrap = True
        tctf.margin_left = tctf.margin_right = Inches(0.18)
        tctf.margin_top = Inches(0.2)
        tctf.margin_bottom = Inches(0.15)

        p = tctf.paragraphs[0]
        p.text = st_title
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = st_col
        p.space_after = Pt(10)

        for pt in st_points:
            pp = tctf.add_paragraph()
            pp.text = f"• {pt}"
            pp.font.size = Pt(8.6)
            pp.font.color.rgb = C_BODY
            pp.space_before = Pt(8)

    pptx_path = os.path.join(PRES_DIR, "Executive_Presentation_Customer_Segmentation_Cohort.pptx")
    prs.save(pptx_path)
    print(f"Minimalist Executive Slide Deck generated successfully: {pptx_path}")
    return pptx_path

if __name__ == "__main__":
    generate_slides()
